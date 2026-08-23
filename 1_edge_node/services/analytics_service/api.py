import io
import os
import sys
import ast
import glob
from typing import Literal
from collections import Counter
from pathlib import Path
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from fastapi import APIRouter
from fastapi.responses import FileResponse, Response, StreamingResponse
from services.database.session import SessionLocal
from services.database.models import InspectionRecord, InspectionObject, AnomalyDetail
from packages.core.config import AppConfig

FILE = Path(__file__).resolve()
ROOT = FILE.parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
cfg = AppConfig.from_yaml(os.path.join(str(ROOT), "config.yaml"))

router = APIRouter(tags=["Analytics"])


"""
23082026 - KHAI - Modify implementation of get_analytics to align with Vietnamese timezone and allow for reuse when exporting
"""
# 23082026-KIET-Query record Analytics theo ngày Việt Nam và task mode để tái sử dụng khi export
def _query_analytics_records(db, date: str, mode: Literal["counting", "defect"]):
    from sqlalchemy import or_
    from sqlalchemy.orm import selectinload
    from services.database.models import InspectionObject

    if date:
        try:
            target_local = datetime.strptime(date, "%Y-%m-%d")
        except ValueError:
            target_local = datetime.utcnow() + timedelta(hours=7)
            target_local = target_local.replace(hour=0, minute=0, second=0, microsecond=0)
    else:
        target_local = datetime.utcnow() + timedelta(hours=7)
        target_local = target_local.replace(hour=0, minute=0, second=0, microsecond=0)

    day_start_utc = target_local - timedelta(hours=7)
    day_end_utc = day_start_utc + timedelta(days=1)

    query = db.query(InspectionRecord).options(
        selectinload(InspectionRecord.objects).selectinload(InspectionObject.anomalies)
    ).filter(
        InspectionRecord.created_at >= day_start_utc,
        InspectionRecord.created_at < day_end_utc,
    )
    if mode == "counting":
        query = query.filter(InspectionRecord.task_type == "detection")
    else:
        query = query.filter(or_(
            InspectionRecord.task_type == "inspection",
            InspectionRecord.task_type.is_(None),
        ))
    return query.order_by(InspectionRecord.created_at.asc()).all()


@router.get("/analytics")
async def get_analytics(
    date: str = None,
    mode: Literal["counting", "defect"] = "defect",
):
    """date: YYYY-MM-DD string in local time (UTC+7). Defaults to today."""

    db = SessionLocal()
    try:
        records = _query_analytics_records(db, date, mode)

        avg_cycle_time = round(
            sum(float(record.latency_ms or 0) for record in records)
            / len(records)
            / 1000,
            3,
        ) if records else 0.0
        visible_hours = set(range(8, 18))
        colors = ["#3b82f6", "#10b981", "#8b5cf6", "#f59e0b", "#f43f5e"]

        if mode == "counting":
            # 23082026-KIET-Tính số object và phân bố class chỉ từ record Detection
            hourly_counts = {hour: 0 for hour in range(24)}
            class_counts = {}
            total_objects = 0
            for record in records:
                object_count = int(record.total_objects or len(record.objects))
                total_objects += object_count
                if record.created_at:
                    local_hour = (record.created_at.hour + 7) % 24
                    hourly_counts[local_hour] += object_count
                for obj in record.objects:
                    class_name = obj.class_name or "unknown"
                    class_counts[class_name] = class_counts.get(class_name, 0) + 1

            active_hours = visible_hours | {
                hour for hour, count in hourly_counts.items() if count > 0
            }
            hourly_counting_data = [
                {"hour": f"{hour:02d}h", "count": hourly_counts[hour]}
                for hour in sorted(active_hours)
            ]
            object_trend_data = [
                {"time": f"{hour:02d}:00", "count": hourly_counts[hour]}
                for hour in sorted(active_hours)
            ]
            class_distribution = [
                {"name": name, "count": count, "color": colors[index % len(colors)]}
                for index, (name, count) in enumerate(
                    sorted(class_counts.items(), key=lambda item: item[1], reverse=True)
                )
            ]
            return {
                "mode": "counting",
                "totalRuns": len(records),
                "totalObjects": total_objects,
                "avgObjectsPerRun": round(total_objects / len(records), 1) if records else 0.0,
                "avgCycleTime": avg_cycle_time,
                "hourlyCountingData": hourly_counting_data,
                "objectTrendData": object_trend_data,
                "classDistribution": class_distribution,
            }

        # 23082026-KIET-Tính Yield và Pareto chỉ từ record Inspection
        hourly_defects = {
            hour: {"ok": 0, "ng": 0}
            for hour in range(24)
        }
        defect_counts = {}
        for record in records:
            local_hour = (
                (record.created_at.hour + 7) % 24
                if record.created_at
                else None
            )
            for obj in record.objects:
                if local_hour is not None:
                    result_key = "ng" if obj.is_ng else "ok"
                    hourly_defects[local_hour][result_key] += 1
                for anomaly in obj.anomalies:
                    defect_name = anomaly.defect_class or "Unknown"
                    defect_counts[defect_name] = defect_counts.get(defect_name, 0) + 1

        active_hours = visible_hours | {
            hour
            for hour, counts in hourly_defects.items()
            if counts["ok"] + counts["ng"] > 0
        }
        hourly_output_data = []
        yield_trend_data = []
        for hour in sorted(active_hours):
            ok_count = hourly_defects[hour]["ok"]
            ng_count = hourly_defects[hour]["ng"]
            total = ok_count + ng_count
            hourly_output_data.append({
                "hour": f"{hour:02d}h",
                "ok": ok_count,
                "ng": ng_count,
            })
            yield_trend_data.append({
                "time": f"{hour:02d}:00",
                "rate": round(ok_count / total * 100, 1) if total else 0.0,
            })

        total_ok = sum(counts["ok"] for counts in hourly_defects.values())
        total_ng = sum(counts["ng"] for counts in hourly_defects.values())
        total_inspected = total_ok + total_ng
        pareto = [
            {"name": name, "count": count, "color": colors[index % len(colors)]}
            for index, (name, count) in enumerate(
                sorted(defect_counts.items(), key=lambda item: item[1], reverse=True)
            )
        ]
        return {
            "mode": "defect",
            "totalInspected": total_inspected,
            "overallYield": round(total_ok / total_inspected * 100, 1) if total_inspected else 100.0,
            "totalNg": total_ng,
            "avgCycleTime": avg_cycle_time,
            "hourlyOutputData": hourly_output_data,
            "yieldTrendData": yield_trend_data,
            "pareto": pareto,
        }
    finally:
        db.close()


"""
23082026 - KHAI - Export analytics report as Excel file consisting of Summary and Raw Data sheets
"""
@router.get("/analytics/export")
async def export_analytics(
    date: str = None,
    mode: Literal["counting", "defect"] = "defect",
):
    """23082026-KIET-Xuất Excel gồm sheet Summary và Raw Data theo Analytics mode."""

    summary = await get_analytics(date=date, mode=mode)
    db = SessionLocal()
    try:
        records = _query_analytics_records(db, date, mode)
        workbook = Workbook()
        summary_sheet = workbook.active
        summary_sheet.title = "Summary"
        raw_sheet = workbook.create_sheet("Raw Data")
        header_fill = PatternFill("solid", fgColor="2563EB")
        header_font = Font(color="FFFFFF", bold=True)

        # 23082026-KIET-Ghi summary và chart source hiện tại vào sheet đầu tiên
        summary_sheet.append(["Vision Analytics", mode.title()])
        summary_sheet.append(["Date", date or "Today"])
        if mode == "counting":
            summary_sheet.append(["Detection Runs", summary["totalRuns"]])
            summary_sheet.append(["Total Objects", summary["totalObjects"]])
            summary_sheet.append(["Average Objects / Run", summary["avgObjectsPerRun"]])
            summary_sheet.append(["Average Cycle Time (s)", summary["avgCycleTime"]])
            summary_sheet.append([])
            summary_sheet.append(["Hour", "Object Count"])
            for row in summary["hourlyCountingData"]:
                summary_sheet.append([row["hour"], row["count"]])
            summary_sheet.append([])
            summary_sheet.append(["Class", "Count"])
            for item in summary["classDistribution"]:
                summary_sheet.append([item["name"], item["count"]])
        else:
            summary_sheet.append(["Total Inspected", summary["totalInspected"]])
            summary_sheet.append(["Overall Yield (%)", summary["overallYield"]])
            summary_sheet.append(["Total NG", summary["totalNg"]])
            summary_sheet.append(["Average Cycle Time (s)", summary["avgCycleTime"]])
            summary_sheet.append([])
            summary_sheet.append(["Hour", "OK", "NG"])
            for row in summary["hourlyOutputData"]:
                summary_sheet.append([row["hour"], row["ok"], row["ng"]])
            summary_sheet.append([])
            summary_sheet.append(["Defect Type", "Count"])
            for item in summary["pareto"]:
                summary_sheet.append([item["name"], item["count"]])

        # 23082026-KIET-Ghi từng object Detection thật vào Raw Data của Counting mode
        if mode == "counting":
            raw_headers = [
                "Record ID", "Record Timestamp", "Created At (UTC+7)", "Camera ID", "Image Path",
                "Object Index", "Class ID", "Class Name", "Confidence",
                "Bounding Box", "Objects In Record", "Latency (ms)",
            ]
            raw_sheet.append(raw_headers)
            for record in records:
                local_timestamp = (
                    record.created_at + timedelta(hours=7)
                    if record.created_at
                    else None
                )
                objects = record.objects or [None]
                for obj in objects:
                    raw_sheet.append([
                        record.id,
                        record.timestamp or "",
                        local_timestamp.strftime("%Y-%m-%d %H:%M:%S") if local_timestamp else "",
                        record.camera_id or "",
                        record.original_image or "",
                        obj.object_index if obj else "",
                        obj.class_id if obj else "",
                        (obj.class_name or "unknown") if obj else "",
                        float(obj.score or 0) if obj else "",
                        (obj.bbox or "") if obj else "",
                        int(record.total_objects or len(record.objects)),
                        float(record.latency_ms or 0),
                    ])
        else:
            # 23082026-KIET-Ghi từng object và anomaly thật vào Raw Data của Defect mode
            raw_headers = [
                "Record ID", "Record Timestamp", "Created At (UTC+7)", "Camera ID", "Image Path",
                "Object Index", "Result", "Object Score", "Object BBox",
                "Overlap Ratio", "Crop Image", "Defect Class", "Similarity",
                "Defect BBox Full", "Defect BBox Crop", "Latency (ms)",
            ]
            raw_sheet.append(raw_headers)
            for record in records:
                local_timestamp = (
                    record.created_at + timedelta(hours=7)
                    if record.created_at
                    else None
                )
                objects = record.objects or [None]
                for obj in objects:
                    anomalies = (obj.anomalies or [None]) if obj else [None]
                    for anomaly in anomalies:
                        raw_sheet.append([
                            record.id,
                            record.timestamp or "",
                            local_timestamp.strftime("%Y-%m-%d %H:%M:%S") if local_timestamp else "",
                            record.camera_id or "",
                            record.original_image or "",
                            obj.object_index if obj else "",
                            ("NG" if obj.is_ng else "OK") if obj else ("NG" if record.ng_detected else "OK"),
                            float(obj.score or 0) if obj else "",
                            (obj.bbox or "") if obj else "",
                            float(obj.overlap_ratio or 0) if obj else "",
                            (obj.crop_image or "") if obj else "",
                            anomaly.defect_class if anomaly else "",
                            float(anomaly.similarity or 0) if anomaly else "",
                            anomaly.bbox_full if anomaly else "",
                            anomaly.bbox_in_object_crop if anomaly else "",
                            float(record.latency_ms or 0),
                        ])

        for sheet in (summary_sheet, raw_sheet):
            sheet.freeze_panes = "A2"
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
            for column_cells in sheet.columns:
                max_length = min(
                    max(len(str(cell.value or "")) for cell in column_cells) + 2,
                    60,
                )
                sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max_length

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        export_date = date or "today"
        filename = f"vision_analytics_{mode}_{export_date}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": f'attachment; filename="{filename}"',
                "Access-Control-Expose-Headers": "Content-Disposition",
            },
        )
    finally:
        db.close()


@router.get("/api/analytics")
def get_legacy_analytics():
    from services.database.models import QCProductPhotoLibrary, InspectionRecord, InspectionObject
    from sqlalchemy import func
    import dateutil.parser
    
    db = SessionLocal()
    try:
        # 1. Defect Pareto from QCProductPhotoLibrary
        counts = db.query(QCProductPhotoLibrary.ErrorDetail, func.count(QCProductPhotoLibrary.ID)).group_by(QCProductPhotoLibrary.ErrorDetail).all()
        
        pareto = []
        colors = ['#f43f5e', '#fb923c', '#facc15', '#38bdf8', '#94a3b8']
        
        for idx, row in enumerate(counts):
            err = row[0] or "Unknown"
            cnt = row[1]
            pareto.append({
                "name": err,
                "count": cnt,
                "color": colors[idx % len(colors)]
            })
            
        pareto = sorted(pareto, key=lambda x: x["count"], reverse=True)
        
        # 2. Production stats from InspectionRecord
        records = db.query(InspectionRecord).all()
        
        total_inspected = len(records)
        total_ng = sum(1 for r in records if r.ng_detected)

        # 11082026 - KIET - Tính tổng object của các record Object Detection.
        total_detected = (
            db.query(func.sum(InspectionRecord.total_objects))
            .filter(InspectionRecord.task_type == "detection")
            .scalar()
            or 0
        )

        # 11082026 - KIET - Thống kê tổng số object theo class đã lưu trong database.
        od_class_rows = (
            db.query(InspectionObject.class_name, func.count(InspectionObject.id))
            .join(InspectionRecord, InspectionObject.record_id == InspectionRecord.id)
            .filter(InspectionRecord.task_type == "detection")
            .group_by(InspectionObject.class_name)
            .all()
        )
        detection_counts_by_class = {
            class_name or "unknown": count
            for class_name, count in od_class_rows
        }
        
        hourly_stats = {}
        
        for r in records:
            try:
                ts_val = r.timestamp or str(r.created_at)
                dt = None
                if "_" in str(ts_val):
                    parts = str(ts_val).split('_')
                    if len(parts) >= 2 and len(parts[0]) == 8 and len(parts[1]) >= 6:
                        d, t = parts[0], parts[1][:6]
                        try:
                            dt = datetime.strptime(f"{d}_{t}", "%Y%m%d_%H%M%S")
                        except ValueError:
                            pass
                if not dt:
                    if isinstance(ts_val, (int, float)):
                        dt = datetime.fromtimestamp(ts_val)
                    else:
                        dt = dateutil.parser.parse(str(ts_val))
                    
                hour_str = f"{dt.hour:02d}h"
                if hour_str not in hourly_stats:
                    hourly_stats[hour_str] = {"ok": 0, "ng": 0}
                    
                if r.ng_detected:
                    hourly_stats[hour_str]["ng"] += 1
                else:
                    hourly_stats[hour_str]["ok"] += 1
                    
            except Exception:
                pass
                
        overall_yield = ((total_inspected - total_ng) / total_inspected * 100) if total_inspected > 0 else 100.0
                
        hourly_output_data = []
        yield_trend_data = []
        
        for hour in sorted(hourly_stats.keys()):
            ok_count = hourly_stats[hour]["ok"]
            ng_count = hourly_stats[hour]["ng"]
            total = ok_count + ng_count
            hourly_output_data.append({
                "hour": hour,
                "ok": ok_count,
                "ng": ng_count
            })
            rate = (ok_count / total * 100) if total > 0 else 100.0
            yield_trend_data.append({
                "time": hour.replace("h", ":00"),
                "rate": round(rate, 1)
            })
            
        if not hourly_output_data:
            hourly_output_data = [{"hour": "08h", "ok": 0, "ng": 0}]
            yield_trend_data = [{"time": "08:00", "rate": 100.0}]

        return {
            "totalInspected": total_inspected,
            # 11082026 - KIET - Bổ sung analytics OD mà không thay response inspection cũ.
            "totalDetected": int(total_detected),
            "detectionCountsByClass": detection_counts_by_class,
            "overallYield": round(overall_yield, 1),
            "totalNg": total_ng,
            "avgCycleTime": 2.45,
            "pareto": pareto,
            "hourlyOutputData": hourly_output_data,
            "yieldTrendData": yield_trend_data
        }
    except Exception as e:
        print(f"Error fetching analytics: {e}")
        return {
            "totalInspected": 0,
            "totalDetected": 0,
            "detectionCountsByClass": {},
            "overallYield": 100.0,
            "totalNg": 0,
            "avgCycleTime": 0.0,
            "pareto": [],
            "hourlyOutputData": [],
            "yieldTrendData": []
        }
    finally:
        db.close()


@router.get("/history")
async def get_history():
    db = SessionLocal()
    try:
        db.commit()
        records = db.query(InspectionRecord).order_by(InspectionRecord.id.asc()).all()
        results = []
        for r in records:
            objs = []
            for o in r.objects:
                anoms = []
                for a in o.anomalies:
                    try:
                        bf = ast.literal_eval(a.bbox_full) if a.bbox_full else []
                    except Exception:
                        bf = []
                    try:
                        bi = ast.literal_eval(a.bbox_in_object_crop) if a.bbox_in_object_crop else []
                    except Exception:
                        bi = []
                    anoms.append({
                        "bbox_full": bf,
                        "bbox_in_object_crop": bi,
                        "cls_label": a.defect_class,
                        "cls_similarity": a.similarity
                    })
                try:
                    obbox = ast.literal_eval(o.bbox) if o.bbox else []
                except Exception:
                    obbox = []
                """
                23082026 - KHAI - Add fields for class counting
                """
                objs.append({
                    "index": o.object_index,
                    "bbox": obbox,
                    "score": o.score,
                    # 11082026 - KIET - Trả metadata OD để frontend hiển thị class counting.
                    "class_id": o.class_id,
                    "class_name": o.class_name,
                    "is_ng": o.is_ng,
                    "overlap_ratio": o.overlap_ratio,
                    "crop": o.crop_image,
                    "anomalies": anoms,
                    "anomaly_count": len(anoms),
                    "anomaly_bboxes_sample": [a["bbox_full"] for a in anoms[:3]]
                })
            
            """
            23082026 - KHAI - Calculate number of objects by class
            """
            # 11082026 - KIET - Tính lại count theo class từ các object đã lưu.
            counts_by_class = {}
            if r.task_type == "detection":
                counts_by_class = dict(
                    Counter(
                        obj["class_name"] or "unknown"
                        for obj in objs
                    )
                )
            
            """
            23082026 - KHAI - Add task type, camera ID fields and fields for class counting
            """
            results.append({
                "timestamp": r.timestamp,
                "image": r.original_image,
                "ng_detected": r.ng_detected,
                "latency_ms": r.latency_ms,
                # 11082026 - KIET - Trả metadata OD nhưng vẫn giữ response inspection cũ.
                "task_type": r.task_type,
                "camera_id": r.camera_id,
                "total_objects": r.total_objects if r.task_type == "detection" else len(objs),
                "counts_by_class": counts_by_class,
                "ng_count": sum(1 for ob in objs if ob["is_ng"]),
                "max_score": max([ob["score"] for ob in objs]) if objs else 0.0,
                "objects": objs
            })
        return {"results": results}
    except Exception as e:
        print(f"Error fetching history: {e}")
        return {"results": []}
    finally:
        db.close()


@router.get("/api/image/{image_name}")
async def get_inspection_image(image_name: str):
    pattern = os.path.join(cfg.CAPTURE_DIR, cfg.PRODUCT_NAME, "sessions", "*", "original", image_name)
    matches = glob.glob(pattern)
    if matches:
        return FileResponse(matches[0])
    return Response(status_code=404)


@router.get("/api/result_image/{image_name}")
async def get_result_image(image_name: str):
    timestamp = image_name.replace("IMG_", "").replace(".jpg", "")
    overall_name = f"OVERALL_{timestamp}.jpg"
    pattern = os.path.join(cfg.CAPTURE_DIR, cfg.PRODUCT_NAME, "sessions", "*", overall_name)
    matches = glob.glob(pattern)
    if matches:
        return FileResponse(matches[0])
    return Response(status_code=404)
